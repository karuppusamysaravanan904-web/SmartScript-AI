"""
HANDWRITE AI - Professional Multi-Format Export Service
Generates authentic Word (.docx), PDF (.pdf), PowerPoint (.pptx), TXT (.txt), MD (.md), and ZIP archives.
"""
import io
import os
import zipfile
from datetime import datetime
from typing import Dict, Any, List

# DOCX
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# PPTX
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor

# PDF
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_docx(doc: Dict[str, Any]) -> bytes:
    """Creates a professional Microsoft Word (.docx) document."""
    d = docx.Document()
    
    # Title
    title = d.add_heading("HANDWRITE AI — Digitized Document", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Metadata Subtitle
    meta_p = d.add_paragraph()
    meta_p.add_run(f"Source File: ").bold = True
    meta_p.add_run(f"{doc.get('filename', 'Unknown')} | ")
    meta_p.add_run(f"Source Script: ").bold = True
    meta_p.add_run(f"{doc.get('source_language', 'Auto Detect')} | ")
    meta_p.add_run(f"Converted Language: ").bold = True
    meta_p.add_run(f"{doc.get('target_language', 'English')}\n")
    meta_p.add_run(f"Digitization Timestamp: ").bold = True
    meta_p.add_run(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | ")
    meta_p.add_run(f"Total Words: ").bold = True
    meta_p.add_run(f"{doc.get('analysis', {}).get('recognized_words', 0)}")
    
    d.add_heading("1. Clean Reconstructed Text", level=1)
    clean_text = doc.get("clean_text", "")
    for para in clean_text.split("\n\n"):
        if para.strip():
            d.add_paragraph(para.strip())
            
    # Translated Output
    target_lang = doc.get("target_language", "English")
    if doc.get("translated_text"):
        d.add_heading(f"2. Converted Multilingual Output ({target_lang})", level=1)
        for para in doc.get("translated_text", "").split("\n\n"):
            if para.strip():
                d.add_paragraph(para.strip())
                
    # Tables if present
    tables = doc.get("tables", [])
    if tables:
        d.add_heading("3. Structured Document Tables", level=1)
        for tbl_data in tables:
            d.add_heading(tbl_data.get("title", "Extracted Table"), level=2)
            headers = tbl_data.get("headers", [])
            rows = tbl_data.get("rows", [])
            
            table = d.add_table(rows=len(rows) + 1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Header Row
            for col_idx, h_text in enumerate(headers):
                cell = table.cell(0, col_idx)
                cell.text = str(h_text)
                cell.paragraphs[0].runs[0].font.bold = True
                
            # Data Rows
            for row_idx, r_data in enumerate(rows):
                for col_idx, val in enumerate(r_data):
                    table.cell(row_idx + 1, col_idx).text = str(val)
                    
            d.add_paragraph() # Spacing
            
    # Uncertainty & Quality Audit Trail
    d.add_heading("4. Handwriting Uncertainty & Forensic Audit", level=1)
    analysis = doc.get("analysis", {})
    stat_p = d.add_paragraph()
    stat_p.add_run(f"High Confidence: {analysis.get('high_confidence', 0)} | ")
    stat_p.add_run(f"Needs Review: {analysis.get('needs_review', 0)} | ")
    stat_p.add_run(f"Unreadable Regions Flagged: {analysis.get('unreadable', 0)}")
    
    # Audit table of uncertain tokens
    tokens = doc.get("tokens", [])
    uncertain_tokens = [t for t in tokens if t.get("is_uncertain") or t.get("is_illegible")]
    if uncertain_tokens:
        d.add_heading("Flagged Low-Confidence Tokens (Anti-Hallucination Guardrails):", level=2)
        audit_tbl = d.add_table(rows=len(uncertain_tokens) + 1, cols=4)
        for c_idx, title_text in enumerate(["Token ID", "Extracted Stroke", "Confidence", "AI Alternative Candidates"]):
            audit_tbl.cell(0, c_idx).text = title_text
            audit_tbl.cell(0, c_idx).paragraphs[0].runs[0].font.bold = True
            
        for r_idx, tok in enumerate(uncertain_tokens):
            audit_tbl.cell(r_idx + 1, 0).text = tok.get("id", "")
            audit_tbl.cell(r_idx + 1, 1).text = tok.get("text", "")
            audit_tbl.cell(r_idx + 1, 2).text = f"{int(tok.get('confidence', 0) * 100)}%"
            audit_tbl.cell(r_idx + 1, 3).text = ", ".join(tok.get("candidates", []))

    buffer = io.BytesIO()
    d.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_pdf(doc: Dict[str, Any]) -> bytes:
    """Creates a formatted PDF document using ReportLab."""
    buffer = io.BytesIO()
    doc_template = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=10
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#4338CA'),
        spaceBefore=14,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=12
    )

    story = []
    # Title
    story.append(Paragraph("HANDWRITE AI — Digitized Document", title_style))
    meta_txt = (
        f"<b>Source:</b> {doc.get('filename', 'Unknown')} &nbsp;|&nbsp; "
        f"<b>Script:</b> {doc.get('source_language', 'Auto')} &nbsp;|&nbsp; "
        f"<b>Output:</b> {doc.get('target_language', 'English')} &nbsp;|&nbsp; "
        f"<b>Words:</b> {doc.get('analysis', {}).get('recognized_words', 0)}"
    )
    story.append(Paragraph(meta_txt, meta_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=12))

    # 1. Recognized Text
    story.append(Paragraph("1. Clean Reconstructed Handwriting Text", h2_style))
    clean_text = doc.get("clean_text", "")
    for p in clean_text.split("\n\n"):
        if p.strip():
            story.append(Paragraph(p.strip().replace("\n", "<br/>"), body_style))
            
    # 2. Translated Text
    if doc.get("translated_text"):
        target_lang = doc.get("target_language", "English")
        story.append(Paragraph(f"2. Converted Multilingual Output ({target_lang})", h2_style))
        for p in doc.get("translated_text", "").split("\n\n"):
            if p.strip():
                story.append(Paragraph(p.strip().replace("\n", "<br/>"), body_style))

    # 3. Tables
    tables = doc.get("tables", [])
    if tables:
        story.append(Paragraph("3. Extracted Document Tables", h2_style))
        for tbl in tables:
            data = [tbl.get("headers", [])] + tbl.get("rows", [])
            t = Table(data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF2FF')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#312E81')),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

    # 4. Uncertainty audit
    story.append(Paragraph("4. Uncertainty & Anti-Hallucination Audit", h2_style))
    analysis = doc.get("analysis", {})
    audit_summary = (
        f"<b>High Confidence Tokens:</b> {analysis.get('high_confidence', 0)} &nbsp;&bull;&nbsp; "
        f"<b>Needs Review:</b> {analysis.get('needs_review', 0)} &nbsp;&bull;&nbsp; "
        f"<b>Unreadable Flagged:</b> {analysis.get('unreadable', 0)}"
    )
    story.append(Paragraph(audit_summary, body_style))

    doc_template.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def generate_pptx(doc: Dict[str, Any]) -> bytes:
    """Creates a multi-slide presentation using python-pptx."""
    prs = Presentation()
    
    # Slide 1: Title Slide
    blank_layout = prs.slide_layouts[6] # Blank
    s1 = prs.slides.add_slide(blank_layout)
    
    tx_box = s1.shapes.add_textbox(PptInches(0.8), PptInches(1.5), PptInches(8.4), PptInches(2.5))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = "HANDWRITE AI"
    p.font.size = PptPt(40)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(79, 70, 229)
    
    p2 = tf.add_paragraph()
    p2.text = "Extreme Bad-Handwriting Digitization & Conversion"
    p2.font.size = PptPt(22)
    p2.font.color.rgb = PptRGBColor(15, 23, 42)
    
    p3 = tf.add_paragraph()
    p3.text = f"Source Document: {doc.get('filename', 'Unknown')} | Target: {doc.get('target_language', 'English')}"
    p3.font.size = PptPt(14)
    p3.font.color.rgb = PptRGBColor(100, 116, 139)
    
    # Slide 2: Recognized Clean Text
    s2 = prs.slides.add_slide(blank_layout)
    t_box = s2.shapes.add_textbox(PptInches(0.8), PptInches(0.6), PptInches(8.4), PptInches(0.8))
    t_tf = t_box.text_frame
    tp = t_tf.paragraphs[0]
    tp.text = "Digitized Handwritten Text"
    tp.font.size = PptPt(26)
    tp.font.bold = True
    tp.font.color.rgb = PptRGBColor(30, 41, 59)
    
    c_box = s2.shapes.add_textbox(PptInches(0.8), PptInches(1.6), PptInches(8.4), PptInches(5.0))
    c_tf = c_box.text_frame
    c_tf.word_wrap = True
    for para in doc.get("clean_text", "").split("\n\n")[:4]:
        if para.strip():
            cp = c_tf.add_paragraph()
            cp.text = para.strip()
            cp.font.size = PptPt(14)
            cp.font.color.rgb = PptRGBColor(51, 65, 85)
            
    # Slide 3: Translated Multilingual Output
    if doc.get("translated_text"):
        s3 = prs.slides.add_slide(blank_layout)
        t_box = s3.shapes.add_textbox(PptInches(0.8), PptInches(0.6), PptInches(8.4), PptInches(0.8))
        tp = t_box.text_frame.paragraphs[0]
        tp.text = f"Converted Output ({doc.get('target_language', 'English')})"
        tp.font.size = PptPt(26)
        tp.font.bold = True
        tp.font.color.rgb = PptRGBColor(16, 185, 129)
        
        c_box = s3.shapes.add_textbox(PptInches(0.8), PptInches(1.6), PptInches(8.4), PptInches(5.0))
        c_tf = c_box.text_frame
        c_tf.word_wrap = True
        for para in doc.get("translated_text", "").split("\n\n")[:4]:
            if para.strip():
                cp = c_tf.add_paragraph()
                cp.text = para.strip()
                cp.font.size = PptPt(14)
                cp.font.color.rgb = PptRGBColor(51, 65, 85)
                
    # Slide 4: Structured Tables if present
    tables = doc.get("tables", [])
    if tables:
        s4 = prs.slides.add_slide(blank_layout)
        t_box = s4.shapes.add_textbox(PptInches(0.8), PptInches(0.6), PptInches(8.4), PptInches(0.8))
        tp = t_box.text_frame.paragraphs[0]
        tp.text = "Extracted Structured Table"
        tp.font.size = PptPt(24)
        tp.font.bold = True
        
        tbl_data = tables[0]
        headers = tbl_data.get("headers", [])
        rows = tbl_data.get("rows", [])
        num_rows = len(rows) + 1
        num_cols = len(headers)
        
        table_shape = s4.shapes.add_table(num_rows, num_cols, PptInches(0.8), PptInches(1.8), PptInches(8.4), PptInches(3.5))
        tbl = table_shape.table
        for c_idx, h in enumerate(headers):
            tbl.cell(0, c_idx).text = str(h)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                tbl.cell(r_idx + 1, c_idx).text = str(val)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_txt(doc: Dict[str, Any]) -> str:
    """Creates formatted plain text."""
    lines = [
        "============================================================",
        "HANDWRITE AI - DIGITIZED HANDWRITTEN DOCUMENT",
        "============================================================",
        f"Filename: {doc.get('filename', 'Unknown')}",
        f"Source Script: {doc.get('source_language', 'Auto Detect')}",
        f"Converted Language: {doc.get('target_language', 'English')}",
        f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "============================================================\n",
        "--- [RECOGNIZED CLEAN TEXT] ---\n",
        doc.get("clean_text", ""),
        "\n--- [CONVERTED TRANSLATION] ---\n",
        doc.get("translated_text", "")
    ]
    tables = doc.get("tables", [])
    if tables:
        lines.append("\n--- [STRUCTURED TABLES] ---\n")
        for tbl in tables:
            lines.append(f"Table: {tbl.get('title', '')}")
            lines.append(" | ".join(tbl.get("headers", [])))
            lines.append("-" * 40)
            for r in tbl.get("rows", []):
                lines.append(" | ".join([str(c) for c in r]))
            lines.append("")
    return "\n".join(lines)

def generate_markdown(doc: Dict[str, Any]) -> str:
    """Creates formatted Markdown with tables and metadata."""
    md = [
        "# HandWrite AI — Digitized Document\n",
        f"- **Source File:** `{doc.get('filename', 'Unknown')}`",
        f"- **Detected Script:** `{doc.get('source_language', 'Auto Detect')}`",
        f"- **Output Language:** `{doc.get('target_language', 'English')}`",
        f"- **Confidence:** `{doc.get('language_detection', {}).get('confidence', 95)}%`\n",
        "## 1. Recognized Text\n",
        doc.get("clean_text", ""),
        f"\n## 2. Converted Translation ({doc.get('target_language', 'English')})\n",
        doc.get("translated_text", "")
    ]
    tables = doc.get("tables", [])
    if tables:
        md.append("\n## 3. Structured Tables\n")
        for tbl in tables:
            headers = tbl.get("headers", [])
            md.append("| " + " | ".join(headers) + " |")
            md.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for r in tbl.get("rows", []):
                md.append("| " + " | ".join([str(c) for c in r]) + " |")
            md.append("")
    return "\n".join(md)

def generate_all_zip(doc: Dict[str, Any]) -> bytes:
    """Bundles DOCX, PDF, PPTX, TXT, MD into a single zip archive."""
    buffer = io.BytesIO()
    doc_id = doc.get("id", "doc")
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"HandWriteAI_{doc_id}.docx", generate_docx(doc))
        z.writestr(f"HandWriteAI_{doc_id}.pdf", generate_pdf(doc))
        z.writestr(f"HandWriteAI_{doc_id}.pptx", generate_pptx(doc))
        z.writestr(f"HandWriteAI_{doc_id}.txt", generate_txt(doc).encode('utf-8'))
        z.writestr(f"HandWriteAI_{doc_id}.md", generate_markdown(doc).encode('utf-8'))
        
    buffer.seek(0)
    return buffer.getvalue()
